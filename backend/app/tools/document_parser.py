from pathlib import Path
from typing import Any, Dict

from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {'.pdf', '.txt', '.docx', '.doc'}


def parse_document(file_path: str) -> Dict[str, Any]:
    path = Path(file_path)
    if not path.exists():
        return {'success': False, 'tool_name': 'document_parser', 'summary': 'Document file was not found.', 'data': {'file_path': file_path}}
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        return {'success': False, 'tool_name': 'document_parser', 'summary': f'Unsupported document type: {extension or "unknown"}.', 'data': {'file_path': file_path, 'extension': extension, 'supported_extensions': sorted(SUPPORTED_EXTENSIONS)}}
    if extension == '.pdf':
        result = _parse_pdf(path)
        if not result['success'] and result.get('data', {}).get('needs_ocr'):
            ocr_result = _ocr_pdf(path)
            if ocr_result['success']:
                return ocr_result
        return result
    if extension == '.txt':
        return _parse_txt(path)
    if extension == '.docx':
        return _parse_docx(path)
    return {'success': False, 'tool_name': 'document_parser', 'summary': '.doc files are not directly parsed. Convert the file to DOCX or PDF and upload it again.', 'data': {'file_path': str(path), 'extension': extension, 'needs_conversion': True}}


def _parse_pdf(path: Path) -> Dict[str, Any]:
    try:
        reader = PdfReader(str(path))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = (page.extract_text() or '').strip()
            except Exception:
                text = ''
            if text:
                pages.append({'page': page_number, 'text': text})
        full_text = '\n\n'.join(f"Page {p['page']}:\n{p['text']}" for p in pages).strip()
        if not full_text:
            return {'success': False, 'tool_name': 'document_parser', 'summary': 'The PDF was opened, but no readable text could be extracted. OCR is required.', 'data': {'file_path': str(path), 'page_count': len(reader.pages), 'text': '', 'pages': [], 'needs_ocr': True}}
        return {'success': True, 'tool_name': 'document_parser', 'summary': f'Text extracted from {len(pages)} PDF page(s).', 'data': {'file_path': str(path), 'file_name': path.name, 'page_count': len(reader.pages), 'text': full_text, 'pages': pages, 'needs_ocr': False, 'extraction_method': 'pdf_text'}}
    except Exception as error:
        return {'success': False, 'tool_name': 'document_parser', 'summary': 'PDF parsing failed.', 'data': {'file_path': str(path), 'error': str(error), 'needs_ocr': False}}


def _ocr_pdf(path: Path) -> Dict[str, Any]:
    try:
        import fitz
        import pytesseract
        from PIL import Image
        import io
        document = fitz.open(str(path))
        pages = []
        for number, page in enumerate(document, start=1):
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes('png')))
            text = pytesseract.image_to_string(image).strip()
            if text:
                pages.append({'page': number, 'text': text})
        document.close()
        full_text = '\n\n'.join(f"Page {p['page']}:\n{p['text']}" for p in pages).strip()
        if not full_text:
            return {'success': False, 'tool_name': 'document_parser', 'summary': 'OCR could not find readable text in the PDF.', 'data': {'file_path': str(path), 'text': '', 'pages': [], 'needs_ocr': True}}
        return {'success': True, 'tool_name': 'document_parser', 'summary': f'OCR extracted text from {len(pages)} PDF page(s).', 'data': {'file_path': str(path), 'file_name': path.name, 'text': full_text, 'pages': pages, 'needs_ocr': False, 'extraction_method': 'ocr'}}
    except Exception as error:
        return {'success': False, 'tool_name': 'document_parser', 'summary': 'OCR fallback was unavailable or failed.', 'data': {'file_path': str(path), 'error': str(error), 'needs_ocr': True}}


def _parse_txt(path: Path) -> Dict[str, Any]:
    try:
        text = path.read_text(encoding='utf-8', errors='replace').strip()
        if not text:
            return {'success': False, 'tool_name': 'document_parser', 'summary': 'The text document is empty.', 'data': {'file_path': str(path), 'text': ''}}
        return {'success': True, 'tool_name': 'document_parser', 'summary': 'Text successfully extracted from document.', 'data': {'file_path': str(path), 'file_name': path.name, 'text': text, 'pages': [{'page': 1, 'text': text}], 'needs_ocr': False, 'extraction_method': 'text'}}
    except Exception as error:
        return {'success': False, 'tool_name': 'document_parser', 'summary': 'Text document parsing failed.', 'data': {'file_path': str(path), 'error': str(error)}}


def _parse_docx(path: Path) -> Dict[str, Any]:
    try:
        from docx import Document
        document = Document(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        table_lines = []
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    table_lines.append(' | '.join(values))
        text = '\n'.join(paragraphs + table_lines).strip()
        if not text:
            return {'success': False, 'tool_name': 'document_parser', 'summary': 'The DOCX file contains no readable text.', 'data': {'file_path': str(path), 'text': ''}}
        return {'success': True, 'tool_name': 'document_parser', 'summary': 'Text successfully extracted from DOCX document.', 'data': {'file_path': str(path), 'file_name': path.name, 'text': text, 'pages': [{'page': 1, 'text': text}], 'needs_ocr': False, 'extraction_method': 'docx'}}
    except Exception as error:
        return {'success': False, 'tool_name': 'document_parser', 'summary': 'DOCX parsing failed. Make sure python-docx is installed.', 'data': {'file_path': str(path), 'error': str(error)}}
